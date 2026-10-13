"""Exports: Companies and People as CSV or Excel, and a Word profile with everything about one company.

Spreadsheet safety: scraped text is untrusted, so any cell that starts with = + - @ (or a tab/CR) is prefixed
with an apostrophe. Otherwise a value such as =HYPERLINK(...) from a web page would run as a formula when a
teammate opens the file (CSV/formula injection).
"""

from __future__ import annotations

import csv
import io
from collections import defaultdict
from datetime import datetime
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .fit import evaluate_fit
from .models import Company, Note, Observation, Person, QualityCheck, Signal, User
from .verify import CHECK_BY_KEY

SOURCE_LABEL = {
    "seed-sheet": "Target sheet",
    "osm": "OpenStreetMap",
    "gleif": "GLEIF",
    "website": "Company website",
    "wikidata": "Wikidata",
    "companies-house": "Companies House",
    "jobs": "Job board",
    "import": "Added by domain",
    "manual": "Edited by the team",
}
SENIORITY_LABEL = {
    "owner": "Owner / founder",
    "c_level": "C-level",
    "partner": "Partner",
    "director": "Director / head",
    "manager": "Manager",
    "staff": "Staff",
}
COMPANY_COLUMNS = [
    ("id", "ID"), ("name", "Company"), ("domain", "Domain"), ("website", "Website"), ("segment", "Segment"),
    ("country", "Country"), ("city", "City"), ("address", "Address"), ("phone", "Phone"),
    ("general_email", "Email"), ("linkedin_url", "LinkedIn"), ("socials", "Other social profiles"),
    ("email_provider", "Email provider"), ("tech", "Website technology"), ("description", "Description"),
    ("tier", "Tier"), ("fit_score", "Fit score"), ("quality_score", "Quality score"), ("status", "Status"),
    ("people", "People"),
    ("decision_makers", "Decision-makers"), ("signals", "Signals"), ("lei", "LEI"),
    ("company_number", "Company number"), ("registry_status", "Registry status"), ("sources", "Sources"),
    ("last_crawled_at", "Website crawled"), ("last_verified_at", "Verified"), ("created_at", "Added"),
]  # fmt: skip
PEOPLE_COLUMNS = [
    ("id", "ID"), ("name", "Name"), ("title", "Title"), ("seniority", "Seniority"), ("email", "Email"),
    ("phone", "Phone"), ("linkedin_url", "LinkedIn"), ("company", "Company"), ("company_domain", "Company domain"),
    ("segment", "Segment"), ("country", "Country"), ("source_url", "Found at"),
]  # fmt: skip


def safe_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def company_rows(s: Session, companies: list[Company]) -> list[dict[str, Any]]:
    ids = [c.id for c in companies]
    if not ids:
        return []
    # .all(): dict(Result) would treat the Result as a mapping (it has .keys()) instead of iterating rows
    people_count: dict[int, int] = dict(
        s.execute(
            select(Person.company_id, func.count()).where(Person.company_id.in_(ids)).group_by(Person.company_id)
        ).all()
    )
    dms: dict[int, list[str]] = defaultdict(list)
    for p in s.scalars(
        select(Person).where(
            Person.company_id.in_(ids), Person.seniority.in_(["owner", "c_level", "partner", "director"])
        )
    ):
        if len(dms[p.company_id]) < 5:
            dms[p.company_id].append(f"{p.name} ({p.title})" if p.title else p.name)
    signals: dict[int, list[str]] = defaultdict(list)
    for sig in s.scalars(select(Signal).where(Signal.company_id.in_(ids)).order_by(Signal.created_at.desc())):
        if len(signals[sig.company_id]) < 8:
            signals[sig.company_id].append(sig.title)
    sources: dict[int, set[str]] = defaultdict(set)
    for cid, src in s.execute(
        select(Observation.entity_id, Observation.source)
        .where(Observation.entity_type == "company", Observation.entity_id.in_(ids))
        .distinct()
    ):
        sources[cid].add(SOURCE_LABEL.get(src, src))
    rows = []
    for c in companies:
        socials = {k: v for k, v in (c.socials or {}).items() if k != "linkedin"}
        rows.append(
            {
                "id": c.id, "name": c.name, "domain": c.domain, "website": c.website, "segment": c.segment,
                "country": c.country, "city": c.city, "address": c.address, "phone": c.phone,
                "general_email": c.general_email, "linkedin_url": c.linkedin_url,
                "socials": " | ".join(f"{k}: {v}" for k, v in sorted(socials.items())),
                "email_provider": c.email_provider, "tech": ", ".join(c.tech or []), "description": c.description,
                "tier": (c.tier or "").title(), "fit_score": c.fit_score, "quality_score": c.quality_score,
                "status": c.status,
                "people": people_count.get(c.id, 0), "decision_makers": "; ".join(dms.get(c.id, [])),
                "signals": "; ".join(signals.get(c.id, [])), "lei": c.lei, "company_number": c.company_number,
                "registry_status": c.registry_status, "sources": ", ".join(sorted(sources.get(c.id, []))),
                "last_crawled_at": c.last_crawled_at, "last_verified_at": c.last_verified_at,
                "created_at": c.created_at,
            }
        )  # fmt: skip
    return rows


def people_rows(s: Session, people: list[tuple[Person, Company]]) -> list[dict[str, Any]]:
    ids = [p.id for p, _ in people]
    found_at: dict[int, str] = {}
    if ids:
        for pid, url in s.execute(
            select(Observation.entity_id, Observation.url).where(
                Observation.entity_type == "person", Observation.entity_id.in_(ids), Observation.field == "name"
            )
        ):
            if url:
                found_at.setdefault(pid, url)
    return [
        {
            "id": p.id,
            "name": p.name,
            "title": p.title,
            "seniority": SENIORITY_LABEL.get(p.seniority or "", ""),
            "email": p.email,
            "phone": p.phone,
            "linkedin_url": p.linkedin_url,
            "company": c.name,
            "company_domain": c.domain,
            "segment": c.segment,
            "country": c.country,
            "source_url": found_at.get(p.id),
        }  # fmt: skip
        for p, c in people
    ]


def to_csv(rows: list[dict[str, Any]], columns: list[tuple[str, str]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([label for _, label in columns])
    for r in rows:
        w.writerow([safe_cell(r.get(k)) for k, _ in columns])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM: Excel opens UTF-8 (Arabic, accents) correctly


HEADER_FILL = PatternFill("solid", fgColor="0F1B2D")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _sheet(ws, rows: list[dict[str, Any]], columns: list[tuple[str, str]]) -> None:
    ws.append([label for _, label in columns])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([safe_cell(r.get(k)) for k, _ in columns])
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, (key, label) in enumerate(columns, start=1):
        sample = [len(str(r.get(key) or "")) for r in rows[:200]]
        width = min(60, max(len(label) + 2, *(sample or [0])) + 2)
        ws.column_dimensions[get_column_letter(i)].width = width


def to_xlsx(
    sheets: list[tuple[str, list[dict[str, Any]], list[tuple[str, str]]]], readme: list[tuple[str, str]]
) -> bytes:
    wb = Workbook()
    first = True
    for title, rows, columns in sheets:
        ws = wb.active if first else wb.create_sheet()
        ws.title = title[:31]
        _sheet(ws, rows, columns)
        first = False
    info = wb.create_sheet("About this export")
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 26, 100
    for k, v in readme:
        info.append([k, safe_cell(v)])
        info.cell(row=info.max_row, column=1).font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def readme_rows(kind: str, count: int, filters: dict[str, Any], user: User) -> list[tuple[str, str]]:
    active = {k: v for k, v in filters.items() if v not in (None, "", [], False)}
    return [
        ("Export", kind),
        ("Rows", str(count)),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Generated by", f"{user.name} <{user.email}>"),
        ("Filters", "; ".join(f"{k}={v}" for k, v in active.items()) or "none"),
        ("Data sources", "Every value is sourced; open a company in Exora to see the source URL of each."),
        ("Use", "Business contact data for B2B outreach. Apply the regional rules in 03 and COMPLIANCE before "
                "contacting anyone (no cold email to KSA personal addresses; EU needs a recorded lawful basis)."),
    ]  # fmt: skip


# ---- Word profile ----------------------------------------------------------------------------------------
def _heading(doc, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x0F, 0x1B, 0x2D)


def _table(doc, headers: list[str], rows: list[list[Any]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)


def company_docx(s: Session, c: Company, user: User) -> bytes:
    people = s.scalars(select(Person).where(Person.company_id == c.id).order_by(Person.seniority, Person.name)).all()
    signals = s.scalars(select(Signal).where(Signal.company_id == c.id).order_by(Signal.created_at.desc())).all()
    checks = s.scalars(
        select(QualityCheck).where(QualityCheck.entity_type == "company", QualityCheck.entity_id == c.id)
    ).all()
    sources = s.scalars(
        select(Observation)
        .where(Observation.entity_type == "company", Observation.entity_id == c.id)
        .order_by(Observation.field, Observation.retrieved_at.desc())
    ).all()
    notes = s.execute(
        select(Note, User).outerjoin(User, User.id == Note.user_id).where(Note.company_id == c.id)
        .order_by(Note.created_at.desc())
    ).all()  # fmt: skip
    fit, reasons = evaluate_fit(s, c)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10.5)
    title = doc.add_heading(c.name, level=0)
    for r in title.runs:
        r.font.color.rgb = RGBColor(0x0F, 0x1B, 0x2D)
    subtitle = " · ".join(x for x in [c.segment, ", ".join(y for y in [c.city, c.country] if y)] if x)
    if subtitle:
        sub = doc.add_paragraph(subtitle)
        sub.runs[0].font.color.rgb = RGBColor(0x5B, 0x65, 0x77)
    meta = doc.add_paragraph(
        f"Company profile generated {datetime.now():%Y-%m-%d %H:%M} by {user.name}. "
        "Every fact below lists its source in the Sources section."
    )
    meta.runs[0].font.size = Pt(8.5)

    _heading(doc, "Summary")
    _table(
        doc,
        ["Field", "Value"],
        [
            ["Website", c.website], ["Phone", c.phone], ["Email", c.general_email], ["Address", c.address],
            ["LinkedIn", c.linkedin_url], ["Email provider", c.email_provider],
            ["Website technology", ", ".join(c.tech or [])], ["Fit score", f"{fit} / 100"],
            ["Quality score", f"{c.quality_score if c.quality_score is not None else '—'} / 100"],
            ["Status", c.status], ["LEI", c.lei], ["Company number", c.company_number],
            ["Registry status", c.registry_status], ["Careers page", c.careers_url],
            ["Website crawled", c.last_crawled_at.strftime("%Y-%m-%d") if c.last_crawled_at else "Not yet"],
        ],
    )  # fmt: skip
    if c.description:
        _heading(doc, "About", 2)
        doc.add_paragraph(c.description)
    others = {k: v for k, v in (c.socials or {}).items()}
    if others:
        _heading(doc, "Social profiles", 2)
        for k, v in sorted(others.items()):
            doc.add_paragraph(f"{k.title()}: {v}", style="List Bullet")

    _heading(doc, "Why this company fits")
    _table(doc, ["Rule", "Points", "Met"], [[r["label"], r["points"], "Yes" if r["met"] else "No"] for r in reasons])

    _heading(doc, f"People ({len(people)})")
    if people:
        _table(
            doc,
            ["Name", "Title", "Seniority", "Email", "Phone"],
            [[p.name, p.title, SENIORITY_LABEL.get(p.seniority or "", ""), p.email, p.phone] for p in people[:200]],
        )
    else:
        doc.add_paragraph("No named people recorded yet.")

    _heading(doc, f"Signals ({len(signals)})")
    if signals:
        _table(
            doc,
            ["Signal", "Type", "Evidence", "Recorded"],
            [[g.title, g.kind, g.url, g.created_at.strftime("%Y-%m-%d")] for g in signals[:100]],
        )
    else:
        doc.add_paragraph("No signals recorded yet.")

    if notes:
        _heading(doc, f"Team notes ({len(notes)})")
        for n, u in notes:
            p = doc.add_paragraph()
            run = p.add_run(f"{n.created_at:%Y-%m-%d} · {u.name if u else 'Unknown'}: ")
            run.bold = True
            p.add_run(n.body)

    _heading(doc, "Data quality")
    _table(
        doc,
        ["Check", "Result", "Detail"],
        [[CHECK_BY_KEY[q.check].label if q.check in CHECK_BY_KEY else q.check, "Pass" if q.passed else "Fail",
          q.detail] for q in checks],
    )  # fmt: skip

    _heading(doc, f"Sources ({len(sources)})")
    _table(
        doc,
        ["Field", "Value", "Source", "Evidence", "Retrieved"],
        [[o.field, (o.value or "")[:120], SOURCE_LABEL.get(o.source, o.source), o.url,
          o.retrieved_at.strftime("%Y-%m-%d")] for o in sources[:400]],
    )  # fmt: skip
    foot = doc.add_paragraph(
        "Business contact data collected from public sources and the company's own website. "
        "Apply the regional outreach rules before contacting anyone."
    )
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.runs[0].font.size = Pt(8)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
