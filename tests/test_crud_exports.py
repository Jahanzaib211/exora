"""CRUD, bulk actions, notes, automations and exports through the HTTP API."""

import io
import zipfile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from prospector.exports import safe_cell, to_csv
from prospector.resolve import CompanyRecord, PersonRecord, upsert_company

from .conftest import login


def seed(db):
    with db.session() as s:
        c, _ = upsert_company(
            s,
            CompanyRecord("osm", "https://osm.org/n/1", "Smith Law", website="smithlaw.co.uk", country="GB",
                          segment="Law", people=[PersonRecord("Jane Doe", "Managing Partner")]),
        )  # fmt: skip
        return c.id


def test_company_crud_with_provenance(app, db):
    c = login(TestClient(app), "sdr@test.io")
    r = c.post("/api/companies", json={"name": "Acme Advisory", "website": "acme-advisory.com", "country": "UK",
                                       "phone": "0161 832 1234", "segment": "CPA"})  # fmt: skip
    assert r.status_code == 201, r.text
    cid = r.json()["id"]
    assert r.json()["country"] == "GB" and r.json()["phone"] == "+441618321234"
    dup = c.post("/api/companies", json={"name": "Acme Advisory Ltd", "website": "https://www.acme-advisory.com/"})
    assert dup.status_code == 409 and dup.json()["detail"]["existing_id"] == cid
    bad = c.patch(f"/api/companies/{cid}", json={"phone": "12"})
    assert bad.status_code == 400
    ok = c.patch(f"/api/companies/{cid}", json={"city": "Leeds", "status": "ready"})
    assert ok.json()["city"] == "Leeds" and ok.json()["status"] == "ready"
    detail = c.get(f"/api/companies/{cid}").json()
    assert any(s["source"] == "manual" and s["field"] == "city" for s in detail["sources"])
    assert detail["fit_reasons"] and detail["quality_score"] is not None
    assert c.delete(f"/api/companies/{cid}?exclude=true").json() == {"deleted": 1}
    assert c.get(f"/api/companies/{cid}").status_code == 404
    again = c.post("/api/companies", json={"name": "Acme Advisory", "website": "acme-advisory.com"})
    assert again.status_code == 201  # manual re-creation is allowed; exclusions only block automated discovery


def test_people_and_notes_crud(app, db):
    cid = seed(db)
    c = login(TestClient(app), "sdr@test.io")
    p = c.post("/api/people", json={"company_id": cid, "name": "Ann Lee", "title": "Head of IT",
                                    "email": "ann.lee@smithlaw.co.uk"})  # fmt: skip
    assert p.status_code == 201 and p.json()["seniority"] == "director"
    pid = p.json()["id"]
    assert c.post("/api/people", json={"company_id": cid, "name": "Ann  Lee"}).status_code == 409
    assert c.patch(f"/api/people/{pid}", json={"email": "not-an-email"}).status_code == 400
    assert c.patch(f"/api/people/{pid}", json={"title": "CFO"}).json()["seniority"] == "c_level"
    assert c.delete(f"/api/people/{pid}").json() == {"deleted": 1}
    n = c.post(f"/api/companies/{cid}/notes", json={"body": "Called reception, ask for Jane."})
    assert n.status_code == 201
    nid = n.json()["id"]
    other = login(TestClient(app), "admin@test.io")
    assert other.patch(f"/api/notes/{nid}", json={"body": "edited by admin"}).status_code == 200
    viewer = login(TestClient(app), "viewer@test.io")
    assert viewer.delete(f"/api/notes/{nid}").status_code == 403
    assert len(c.get(f"/api/companies/{cid}/notes").json()) == 1
    assert c.delete(f"/api/notes/{nid}").json() == {"deleted": 1}


def test_bulk_actions(app, db):
    cid = seed(db)
    c = login(TestClient(app), "sdr@test.io")
    assert c.post("/api/companies/bulk", json={"ids": [cid], "action": "status", "status": "ready"}).json() == {
        "updated": 1
    }
    enrich = c.post("/api/companies/bulk", json={"ids": [cid], "action": "enrich"}).json()
    assert enrich["queued"] == 1 and len(enrich["runs"]) == 1
    runs = c.get("/api/runs").json()
    assert runs[0]["source"] == "website" and runs[0]["params"]["companies"] == 1
    assert c.post(f"/api/runs/{runs[0]['id']}/cancel").json()["status"] == "cancelled"
    assert c.post(f"/api/runs/{runs[0]['id']}/cancel").status_code == 409
    assert c.post("/api/companies/bulk", json={"ids": [cid], "action": "delete"}).json() == {"deleted": 1}


def test_automation_crud(app, db):
    c = login(TestClient(app), "sdr@test.io")
    bad = c.post("/api/automations", json={"name": "x", "source": "osm", "params": {"segment": "Law"},
                                           "rotate_field": "city", "rotate_values": []})  # fmt: skip
    assert bad.status_code == 400
    r = c.post("/api/automations", json={"name": "Leeds law", "source": "osm", "params": {"segment": "Law"},
                                         "rotate_field": "city", "rotate_values": ["Leeds", "York"],
                                         "cadence_minutes": 30})  # fmt: skip
    assert r.status_code == 201, r.text
    a = r.json()
    assert a["next_value"] == "Leeds" and a["enabled"]
    run = c.post(f"/api/automations/{a['id']}/run").json()
    assert run["params"]["city"] == "Leeds" and run["automation"] == "Leeds law"
    assert c.post(f"/api/automations/{a['id']}/run").status_code == 409  # already queued
    upd = c.patch(f"/api/automations/{a['id']}", json={"enabled": False, "cadence_minutes": 120}).json()
    assert upd["enabled"] is False and upd["next_run_at"] is None
    assert c.delete(f"/api/automations/{a['id']}").json() == {"deleted": 1}
    eng = c.get("/api/engine").json()
    assert eng["queue"]["queued"] >= 1 and eng["series"]


def test_exports(app, db):
    seed(db)
    c = login(TestClient(app), "sdr@test.io")
    csv_r = c.get("/api/export/companies", params={"format": "csv"})
    assert csv_r.status_code == 200 and "attachment" in csv_r.headers["content-disposition"]
    assert "Smith Law" in csv_r.content.decode("utf-8-sig")
    x = c.get("/api/export/companies", params={"format": "xlsx", "segment": "Law"})
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Companies", "People", "About this export"]
    assert wb["People"]["B2"].value == "Jane Doe"
    cid = c.get("/api/companies").json()["items"][0]["id"]
    d = c.get(f"/api/companies/{cid}/profile.docx")
    assert d.status_code == 200
    with zipfile.ZipFile(io.BytesIO(d.content)) as z:
        xml = z.read("word/document.xml").decode()
    assert "Smith Law" in xml and "Why this company fits" in xml and "Sources" in xml
    viewer = login(TestClient(app), "viewer@test.io")
    assert viewer.get("/api/export/companies").status_code == 403  # bulk export needs the sdr role


def test_formula_injection_guard():
    assert safe_cell('=HYPERLINK("x")') == '\'=HYPERLINK("x")'
    assert safe_cell("+44 161") == "'+44 161"
    assert safe_cell("Normal") == "Normal" and safe_cell(None) == ""
    assert b"'=1+1" in to_csv([{"a": "=1+1"}], [("a", "A")])
